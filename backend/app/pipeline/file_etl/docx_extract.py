"""DOCX extraction — paragraphs, tables, embedded images."""

from __future__ import annotations

import io
from typing import TYPE_CHECKING

from .types import ExtractionStats, ExtractedDocument, ImageNote, TableBlock
from .vision import MAX_IMAGES_PER_FILE, describe_image

if TYPE_CHECKING:
    from collections.abc import Callable


def extract_docx(
    data: bytes,
    filename: str,
    *,
    on_progress: Callable[[str], None] | None = None,
) -> ExtractedDocument:
    from docx import Document

    doc = Document(io.BytesIO(data))
    stats = ExtractionStats()

    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)
    if on_progress:
        on_progress(f"Extracted {len(paragraphs)} paragraphs from DOCX")

    tables: list[TableBlock] = []
    for ti, tbl in enumerate(doc.tables):
        rows = []
        for row in tbl.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        if not rows:
            continue
        header = rows[0]
        md_header = "| " + " | ".join(header) + " |"
        md_sep = "| " + " | ".join("---" for _ in header) + " |"
        md_rows = "\n".join("| " + " | ".join(r) + " |" for r in rows[1:])
        tables.append(TableBlock(label=f"Table {ti + 1}", content=f"{md_header}\n{md_sep}\n{md_rows}"))
    stats.tables_found = len(tables)
    if tables and on_progress:
        on_progress(f"Found {len(tables)} table(s) in DOCX")

    image_notes: list[ImageNote] = []
    image_parts = [rel for rel in doc.part.rels.values() if "image" in (rel.reltype or "")]
    img_count = 0
    for rel in image_parts:
        if img_count >= MAX_IMAGES_PER_FILE:
            stats.images_skipped += 1
            continue
        try:
            img_bytes = rel.target_part.blob
            if len(img_bytes) < 2000:
                continue
            ct = rel.target_part.content_type or "image/png"
            desc = describe_image(img_bytes, mime=ct, context=f"From {filename}")
            image_notes.append(ImageNote(label=f"embedded_image_{img_count + 1}", description=desc))
            stats.images_processed += 1
            img_count += 1
        except Exception:
            stats.images_skipped += 1

    if image_notes and on_progress:
        on_progress(f"Described {len(image_notes)} embedded image(s)")

    return ExtractedDocument(
        source_name=filename,
        format="docx",
        text=text,
        tables=tables,
        image_notes=image_notes,
        stats=stats,
    )
