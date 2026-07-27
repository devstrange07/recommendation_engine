"""XLSX extraction — each sheet as a markdown table."""

from __future__ import annotations

import io
from typing import TYPE_CHECKING

from .types import ExtractionStats, ExtractedDocument, TableBlock

if TYPE_CHECKING:
    from collections.abc import Callable


def extract_xlsx(
    data: bytes,
    filename: str,
    *,
    on_progress: Callable[[str], None] | None = None,
) -> ExtractedDocument:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    stats = ExtractionStats(sheets=len(wb.sheetnames))
    if on_progress:
        on_progress(f"Detected XLSX ({len(wb.sheetnames)} sheet(s))")

    tables: list[TableBlock] = []
    text_parts: list[str] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        str_rows = [[str(c) if c is not None else "" for c in row] for row in rows]
        header = str_rows[0]
        md_header = "| " + " | ".join(header) + " |"
        md_sep = "| " + " | ".join("---" for _ in header) + " |"
        md_rows = "\n".join("| " + " | ".join(r) + " |" for r in str_rows[1:])
        md = f"{md_header}\n{md_sep}\n{md_rows}"
        tables.append(TableBlock(label=f"Sheet: {sheet_name}", content=md))

        text_parts.append(f"[Sheet: {sheet_name}]\n" + "\n".join(
            ", ".join(r) for r in str_rows[:5]
        ))

    wb.close()
    stats.tables_found = len(tables)
    if on_progress:
        on_progress(f"Extracted {len(tables)} sheet(s) as tables")

    return ExtractedDocument(
        source_name=filename,
        format="xlsx",
        text="\n\n".join(text_parts),
        tables=tables,
        stats=stats,
    )
