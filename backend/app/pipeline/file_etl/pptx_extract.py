"""PPTX extraction — slide text, tables, images."""

from __future__ import annotations

import io
from typing import TYPE_CHECKING

from .types import ExtractionStats, ExtractedDocument, ImageNote, TableBlock
from .vision import MAX_IMAGES_PER_FILE, describe_image

if TYPE_CHECKING:
    from collections.abc import Callable


def extract_pptx(
    data: bytes,
    filename: str,
    *,
    on_progress: Callable[[str], None] | None = None,
) -> ExtractedDocument:
    from pptx import Presentation
    from pptx.shapes.picture import Picture
    from pptx.shapes.graphfrm import GraphicFrame

    prs = Presentation(io.BytesIO(data))
    stats = ExtractionStats(pages=len(prs.slides))
    if on_progress:
        on_progress(f"Detected PPTX ({len(prs.slides)} slides)")

    text_parts: list[str] = []
    tables: list[TableBlock] = []
    images_to_describe: list[tuple[bytes, str, str]] = []

    for si, slide in enumerate(prs.slides):
        slide_texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        slide_texts.append(t)

            if isinstance(shape, GraphicFrame) and shape.has_table:
                tbl = shape.table
                rows = []
                for row in tbl.rows:
                    rows.append([cell.text.strip() for cell in row.cells])
                if rows:
                    header = rows[0]
                    md_header = "| " + " | ".join(header) + " |"
                    md_sep = "| " + " | ".join("---" for _ in header) + " |"
                    md_rows = "\n".join("| " + " | ".join(r) + " |" for r in rows[1:])
                    tables.append(TableBlock(
                        label=f"Slide {si + 1}",
                        content=f"{md_header}\n{md_sep}\n{md_rows}",
                    ))

            if isinstance(shape, Picture) and len(images_to_describe) < MAX_IMAGES_PER_FILE:
                try:
                    blob = shape.image.blob
                    ct = shape.image.content_type or "image/png"
                    if len(blob) >= 2000:
                        images_to_describe.append((blob, f"slide_{si + 1}", ct))
                except Exception:
                    pass

        if slide_texts:
            text_parts.append(f"[Slide {si + 1}]\n" + "\n".join(slide_texts))

    text = "\n\n".join(text_parts)
    if on_progress:
        on_progress(f"Extracted text from {stats.pages} slides")

    stats.tables_found = len(tables)
    if tables and on_progress:
        on_progress(f"Found {len(tables)} table(s) in slides")

    image_notes: list[ImageNote] = []
    if images_to_describe and on_progress:
        on_progress(f"Describing {len(images_to_describe)} image(s) via vision…")
    for img_bytes, label, mime in images_to_describe:
        try:
            desc = describe_image(img_bytes, mime=mime, context=f"From {filename}, {label}")
            image_notes.append(ImageNote(label=label, description=desc))
            stats.images_processed += 1
        except Exception:
            stats.images_skipped += 1

    return ExtractedDocument(
        source_name=filename,
        format="pptx",
        text=text,
        tables=tables,
        image_notes=image_notes,
        stats=stats,
    )
